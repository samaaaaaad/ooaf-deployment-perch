"""Build the architecture-only report draft with editable prose and UML figures."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

B=Path(__file__).resolve().parent
import json
SOURCE_COMMIT=json.loads((B/'class-inventory.json').read_text())['commit']
doc=Document()
sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(20);sec.bottom_margin=Mm(20);sec.left_margin=Mm(23);sec.right_margin=Mm(23)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Title'].font.size=Pt(21)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 2'].font.size=Pt(13)
doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].font.bold=False
# The bundled Word base may carry inherited title/paragraph borders.
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.core_properties.title='Perch Project Class Architecture'
doc.core_properties.subject='ELEC5619 final report Section 2'
doc.core_properties.author=''
header=sec.header.paragraphs[0]
header.text='ELEC5619  |  Perch  |  Class architecture'
header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Page ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(s):doc.add_paragraph(s)
def heading(s):doc.add_heading(s,level=2)
def landscape():
    s=doc.add_section(WD_SECTION_START.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE;s.page_width=Mm(420);s.page_height=Mm(297)
    s.left_margin=Mm(15);s.right_margin=Mm(15);s.top_margin=Mm(15);s.bottom_margin=Mm(15)
    return s
def portrait():
    s=doc.add_section(WD_SECTION_START.NEW_PAGE)
    s.orientation=WD_ORIENT.PORTRAIT;s.page_width=Mm(210);s.page_height=Mm(297)
    s.left_margin=Mm(23);s.right_margin=Mm(23);s.top_margin=Mm(20);s.bottom_margin=Mm(20)
def figure(filename,caption,alt):
    par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after=Pt(3)
    shape=par.add_run().add_picture(str(B/filename),width=Mm(350))
    shape._inline.docPr.set('descr',alt)
    doc.add_paragraph(caption,'Caption')

doc.add_paragraph('2 Project Architecture','Title')
p('Perch supports ticket sales for temporary and mobile event businesses through a React interface, a Spring Boot application and a local H2 database (PostgreSQL in the production profile). The architecture separates HTTP handling, business rules and persistence. Organisers configure reusable seating layouts and publish events; customers reserve event seats, complete checkout and retrieve electronic tickets.')
p(f'This class-diagram section describes main at source commit {SOURCE_COMMIT}, inspected on 29 October 2026. Its two complementary views separate the persistent domain from the application services. They use the class compartments, visibility and numeric multiplicities taught in Lecture 3, pages 10–17. The separately required use case diagram is outside this section.')
heading('2.1 Class Diagram')
p('Figure 2.1a shows the ten persistent domain classes. User represents customers, organisers, pending organisers and administrators through a UserRole value. OrganiserRequest and StudentVerificationRequest retain application outcomes; approved student eligibility is recorded on User rather than represented by a subclass. A user may organise multiple events, own venues and floor plans, and place bookings.')
p('FloorPlan stores a reusable layout and its Seat records. Event optionally references a floor plan while in draft; publication requires a configured plan and creates EventSeat inventory. Each EventSeat stores its own price, availability and hold state. Seat and EventSeat have no direct object association. Event stores category, venueName and city as text but has no Venue object association. These display values do not introduce another persistent relationship.')
p('A Booking belongs to one customer and one event. Its seatIds record the checkout selection; sold EventSeat objects link to the booking on confirmation. Each Ticket belongs to one booking and exactly one event seat, with a unique seat reference and ticket code. Optional multiplicities allow for unsold seats, draft events and pending or cancelled bookings. A confirmed purchase issues one ticket per sold seat.')
p('Figure 2.1b shows the business classes. EventService manages events and floor plans; VenueService provides organiser-scoped venue reads; venue creation, editing and deletion are absent. UserService handles credentials and registration, while the two request services support administrator review. AdminService provides revenue and popularity statistics. Controllers invoke these services, which use Spring Data repository interfaces to persist the domain objects. DTOs carry response and request data across the HTTP boundary; the figures omit repetitive accessors, DTO classes and individual controller/repository classes.')
p('SeatHoldService uses EventSeat versioning to detect competing holds, and ExpiredSeatHoldReleaseJob frees expired reservations. BookingService coordinates checkout with PaymentProvider, TicketService and BookingEmailService. MockPaymentProvider confirms locally; StripePaymentProvider creates a payment intent and confirmation follows a verified webhook. BookingRepository applies a write lock when looking up the payment reference. Ticket issuance shares the confirmation transaction. Email is best effort and requires mail configuration.')

landscape()
figure('domain-classes.png','Figure 2.1a  UML class diagram — domain structure. Selected members and association multiplicities reflect the inspected branch.','Ten persistent classes and fourteen associations, including optional floor-plan ownership and the unique ticket-to-event-seat association.')
landscape()
figure('application-classes.png','Figure 2.1b  UML class diagram — application structure. Dashed open arrows denote dependencies; hollow triangles denote interface realisation.','All eleven service and scheduled-job classes and the PaymentProvider interface with its mock and Stripe implementations.')
target=B/'Project Class Architecture.docx'
doc.save(target)
check=Document(target)
assert len(check.inline_shapes)==2 and len(check.sections)==3
assert all(s._inline.docPr.get('descr') for s in check.inline_shapes)
assert not any(p.text.startswith('2.2') for p in check.paragraphs)
assert SOURCE_COMMIT in '\n'.join(p.text for p in check.paragraphs)
print('PASS: saved class-only Word excerpt with two figures, three sections and the source commit.')
